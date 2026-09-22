import argparse
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def load_data(file_path: Path) -> pd.DataFrame:
    if file_path.suffix.lower() in {".xlsx", ".xls"}:
        return pd.read_excel(file_path)
    return pd.read_csv(file_path)


def analyze(file_path: Path) -> dict:
    df = load_data(file_path)
    df["transaction_date"] = pd.to_datetime(df["transaction_date"], format="%m/%d/%y")
    df["transaction_time"] = pd.to_datetime(df["transaction_time"], format="%H:%M:%S")
    df["revenue"] = df["transaction_qty"] * df["unit_price"]
    df["hour"] = df["transaction_time"].dt.hour

    return {
        "df": df,
        "total_revenue": df["revenue"].sum(),
        "transactions": len(df),
        "items_sold": int(df["transaction_qty"].sum()),
        "avg_value": float(np.mean(df["revenue"])),
        "median_value": float(np.median(df["revenue"])),
        "date_range": (df["transaction_date"].min(), df["transaction_date"].max()),
        "by_category": df.groupby("product_category")["revenue"].sum().sort_values(ascending=False),
        "by_store": df.groupby("store_location")["revenue"].sum().sort_values(ascending=False),
        "by_hour": df.groupby("hour")["revenue"].sum().sort_index(),
        "top_products": df.groupby("product_detail")["revenue"].sum().sort_values(ascending=False).head(10),
    }


def save_chart(fig: plt.Figure, path: Path) -> None:
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close(fig)


def build_charts(r: dict, output_dir: Path) -> dict[str, Path]:
    sns.set_theme(style="whitegrid")
    paths = {}

    # Revenue by category
    data = r["by_category"].reset_index()
    data.columns = pd.Index(["product_category", "revenue"])
    fig, ax = plt.subplots(figsize=(8, 5))
    sns.barplot(data=data, x="revenue", y="product_category", hue="product_category", legend=False, palette="viridis", ax=ax)
    ax.set_title("Revenue by Product Category")
    ax.set_xlabel("Revenue ($)")
    ax.set_ylabel("")
    fig.tight_layout()
    paths["category"] = output_dir / "chart_category.png"
    save_chart(fig, paths["category"])

    # Revenue by store
    data = r["by_store"].reset_index()
    data.columns = pd.Index(["store_location", "revenue"])
    fig, ax = plt.subplots(figsize=(8, 4))
    sns.barplot(data=data, x="revenue", y="store_location", hue="store_location", legend=False, palette="mako", ax=ax)
    ax.set_title("Revenue by Store")
    ax.set_xlabel("Revenue ($)")
    ax.set_ylabel("")
    fig.tight_layout()
    paths["store"] = output_dir / "chart_store.png"
    save_chart(fig, paths["store"])

    # Revenue by hour
    hourly = r["by_hour"].reset_index()
    hourly.columns = pd.Index(["hour", "revenue"])
    fig, ax = plt.subplots(figsize=(8, 4))
    sns.lineplot(data=hourly, x="hour", y="revenue", marker="o", ax=ax)
    ax.set_title("Revenue by Transaction Hour")
    ax.set_xlabel("Hour of Day")
    ax.set_ylabel("Revenue ($)")
    ax.set_xticks(hourly["hour"])
    fig.tight_layout()
    paths["hourly"] = output_dir / "chart_hourly.png"
    save_chart(fig, paths["hourly"])

    return paths


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (key, val) in enumerate(rows):
        table.rows[i].cells[0].text = key
        table.rows[i].cells[1].text = val
        for cell in table.rows[i].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(11)


def add_series_table(doc: Document, series: pd.Series, col1: str, col2: str) -> None:
    table = doc.add_table(rows=1 + len(series), cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = col1
    hdr[1].text = col2
    for cell in hdr:
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else None
    for i, (idx, val) in enumerate(series.items(), start=1):
        table.rows[i].cells[0].text = str(idx)
        table.rows[i].cells[1].text = f"${val:,.2f}"


def generate_word_report(r: dict, charts: dict[str, Path], output_path: Path) -> None:
    doc = Document()

    # Title
    title = doc.add_heading("Coffee Shop Sales Analysis Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    start, end = r["date_range"]
    doc.add_paragraph(f"Report Period: {start:%B %d, %Y} – {end:%B %d, %Y}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # Summary
    add_heading(doc, "1. Executive Summary")
    add_kv_table(doc, [
        ("Date Range", f"{start:%Y-%m-%d} to {end:%Y-%m-%d}"),
        ("Total Transactions", f"{r['transactions']:,}"),
        ("Total Items Sold", f"{r['items_sold']:,}"),
        ("Total Revenue", f"${r['total_revenue']:,.2f}"),
        ("Average Transaction Value", f"${r['avg_value']:,.2f}"),
        ("Median Transaction Value", f"${r['median_value']:,.2f}"),
    ])
    doc.add_paragraph()

    # Revenue by Category
    add_heading(doc, "2. Revenue by Product Category")
    add_series_table(doc, r["by_category"], "Product Category", "Revenue")
    doc.add_paragraph()
    doc.add_picture(str(charts["category"]), width=Inches(5.5))
    doc.add_paragraph()

    # Revenue by Store
    add_heading(doc, "3. Revenue by Store")
    add_series_table(doc, r["by_store"], "Store Location", "Revenue")
    doc.add_paragraph()
    doc.add_picture(str(charts["store"]), width=Inches(5.5))
    doc.add_paragraph()

    # Revenue by Hour
    add_heading(doc, "4. Revenue by Transaction Hour")
    doc.add_paragraph("The chart below shows revenue distribution across hours of the day.")
    doc.add_picture(str(charts["hourly"]), width=Inches(5.5))
    doc.add_paragraph()

    # Top Products
    add_heading(doc, "5. Top 10 Products by Revenue")
    add_series_table(doc, r["top_products"], "Product", "Revenue")
    doc.add_paragraph()

    # About
    add_heading(doc, "6. About This Report")
    doc.add_paragraph(
        "This report was generated using Python with the following libraries: "
        "pandas (data processing), NumPy (statistical calculations), "
        "Matplotlib and Seaborn (data visualisation), and python-docx (Word report generation). "
        "Revenue is calculated as transaction_qty × unit_price."
    )

    doc.save(output_path)
    print(f"Report saved to: {output_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate Word report for coffee sales.")
    parser.add_argument("file", nargs="?", type=Path, default=Path("Coffee Shop Sales.csv"))
    parser.add_argument("--output", type=Path, default=Path("Coffee_Sales_Report.docx"))
    parser.add_argument("--chart-dir", type=Path, default=Path("analysis_outputs"))
    args = parser.parse_args()

    args.chart_dir.mkdir(parents=True, exist_ok=True)
    r = analyze(args.file)
    charts = build_charts(r, args.chart_dir)
    generate_word_report(r, charts, args.output)


if __name__ == "__main__":
    main()
